# -*- coding: utf-8 -*-
"""Genera cuadro de prensa actualizado con Int 2 LDR-CIE corregido."""

# ── DATOS ACTUALIZADOS ────────────────────────────────────────────────────────
# Int 2 LDR-CIE: Postobon 470.902t (antes 384t) + Bavaria 264t = 734.902t
# Cambios vs version anterior: solo Postobon Marzo +86.902t

# CIE-LDR (por tipo de carga, por mes)
papel_mar  = 681.36 + 510 + 541.31        # Int1+2+3 CIE-LDR GNL
papel_abr  = 135.12                        # Int4
papel_may  = 169.18 + 475.63              # Int5+6

varilla_mar = 327.36 + 435.22 + 280.45   # Int1+2+3
varilla_abr = 634.07
varilla_may = 447.83 + 429.78

palank_mar  = 167.04 + 222.09 + 143.11
palank_abr  = 323.56
palank_may  = 136.44

estibas_mar = 8.0

# LDR-CIE Postobon (ACTUALIZADO Int2: 470.902 en vez de 384)
pos_mar = 712.8 + 470.902    # Int1=712.8, Int2=470.902
pos_abr = 611.107 + 414.98  # Int3+4
pos_may = 395.99 + 477.508 + 230.68 + 74.83  # Int5+6 LDR + Int5+6 CIE (Postobon en CIE-LDR)

# LDR-CIE Bavaria
bav_mar = 204 + 264           # Int1+2
bav_abr = 270.72              # Int4
bav_may = 205.12 + 169.2     # Int5+6

# Totales por mes
tot_mar = papel_mar + varilla_mar + palank_mar + estibas_mar + pos_mar + bav_mar
tot_abr = papel_abr + varilla_abr + palank_abr + pos_abr + bav_abr
tot_may = papel_may + varilla_may + palank_may + pos_may + bav_may

def f(v): return f"{v:,.1f}".replace(",","X").replace(".",",").replace("X",".")

FILAS = [
    ("Papel",               "GNL",           f(papel_mar),  f(papel_abr),  f(papel_may),  f(papel_mar+papel_abr+papel_may)),
    ("Acero - Varilla",     "Ternium Steel",  f(varilla_mar),f(varilla_abr),f(varilla_may),f(varilla_mar+varilla_abr+varilla_may)),
    ("Acero - Palanquilla", "Diaco",          f(palank_mar), f(palank_abr), f(palank_may), f(palank_mar+palank_abr+palank_may)),
    ("Bebidas",             "Postobon",       f(pos_mar),    f(pos_abr),    f(pos_may),    f(pos_mar+pos_abr+pos_may)),
    ("Bebidas",             "Bavaria",        f(bav_mar),    f(bav_abr),    f(bav_may),    f(bav_mar+bav_abr+bav_may)),
    ("Estibas",             "Transferport",   f(estibas_mar),"--",          "--",          f(estibas_mar)),
]
TOTAL_ROW = ("TOTAL", "", f(tot_mar), f(tot_abr), f(tot_may), f(tot_mar+tot_abr+tot_may))

print("=== VERIFICACION ===")
for row in FILAS:
    print(row)
print("TOTAL:", TOTAL_ROW)

from fpdf import FPDF
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CABECERAS = ["Tipo de Carga", "Cliente", "Marzo 2026", "Abril 2026", "Mayo 2026", "TOTAL"]
ANCHOS_PDF = [40, 34, 27, 27, 27, 27]
NAVY  = (26, 58, 92)
LBLUE = (221, 232, 244)
LGRAY = (245, 247, 250)
WHITE = (255, 255, 255)
GRAY  = (85, 101, 122)

# ── PDF ───────────────────────────────────────────────────────────────────────
class PDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica","I",8)
        self.set_text_color(*GRAY)
        self.cell(0,5,"Corredor Ferroviario Cienaga - La Dorada  |  Periodo Marzo - Mayo 2026",align="C")

pdf = PDF(orientation="L", unit="mm", format="A4")
pdf.set_auto_page_break(auto=False)
pdf.add_page()
pdf.set_margins(14,14,14)

pdf.set_font("Helvetica","I",9)
pdf.set_text_color(*GRAY)
pdf.set_xy(14,14)
pdf.cell(0,6,"Bogota, 28 de mayo de 2026",align="R")
pdf.ln(8)
pdf.set_font("Helvetica","B",17)
pdf.set_text_color(*NAVY)
pdf.cell(0,8,"CARGA TRANSPORTADA",align="C"); pdf.ln(9)
pdf.set_font("Helvetica","B",12)
pdf.cell(0,6,"Corredor Ferreo Cienaga <-> La Dorada",align="C"); pdf.ln(7)
pdf.set_font("Helvetica","I",10)
pdf.set_text_color(*GRAY)
pdf.cell(0,5,"Periodo: Marzo - Mayo 2026  |  Unidad: Toneladas metricas (t)",align="C"); pdf.ln(4)
pdf.set_draw_color(*NAVY); pdf.set_line_width(0.6)
pdf.line(14,pdf.get_y(),pdf.w-14,pdf.get_y()); pdf.ln(6)

LINE_H=9
y_table=pdf.get_y()

def draw_row(cells,bg,txt,bold=False,fs=9.5,lh=LINE_H):
    pdf.set_fill_color(*bg); pdf.set_text_color(*txt)
    pdf.set_font("Helvetica","B" if bold else "",fs)
    for i,(t,w) in enumerate(zip(cells,ANCHOS_PDF)):
        pdf.cell(w,lh,t,border=0,align="R" if i>=2 else "L",fill=True)
    pdf.ln(lh)

draw_row(CABECERAS,NAVY,WHITE,bold=True,fs=9.5,lh=10)
pdf.set_draw_color(*NAVY); pdf.set_line_width(0.3)
pdf.line(14,pdf.get_y(),pdf.w-14,pdf.get_y())
for idx,row in enumerate(FILAS):
    draw_row(row,LGRAY if idx%2==1 else WHITE,(30,30,30),fs=9.5)
pdf.set_line_width(0.4); pdf.line(14,pdf.get_y(),pdf.w-14,pdf.get_y())
draw_row(TOTAL_ROW,LBLUE,NAVY,bold=True,fs=10,lh=10)
y_end=pdf.get_y()
pdf.set_line_width(0.5); pdf.rect(14,y_table,sum(ANCHOS_PDF),y_end-y_table)
pdf.ln(5)
pdf.set_font("Helvetica","I",8.5); pdf.set_text_color(*GRAY)
pdf.multi_cell(0,4.5,
    "Nota: Corresponde a 12 trenes intermodales operados en el periodo "
    "(6 en la direccion Cienaga -> La Dorada y 6 en La Dorada -> Cienaga). "
    "Las cifras incluyen toda la carga movilizada en ambos sentidos del corredor.",align="L")

OUT_PDF = r"C:\Users\jony9\Downloads\Jony 11-25\Carga_Transportada_Mar-May2026.pdf"
pdf.output(OUT_PDF)
print(f"PDF OK: {OUT_PDF}")

# ── WORD ─────────────────────────────────────────────────────────────────────
NAVY_HEX="1A3A5C"; LBLUE_HEX="DDE8F4"; LGRAY_HEX="F5F7FA"; WHITE_HEX="FFFFFF"
NAVY_RGB=RGBColor(0x1a,0x3a,0x5c); WHITE_RGB=RGBColor(0xFF,0xFF,0xFF)
BLACK_RGB=RGBColor(0,0,0); GRAY_RGB=RGBColor(0x55,0x65,0x7A)
ANCHOS_W=[4.2,3.6,2.8,2.8,2.8,2.8]

def set_bg(cell,hex_color):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hex_color)
    for old in tcPr.findall(qn('w:shd')): tcPr.remove(old)
    tcPr.append(shd)

def set_borders(cell,color="D0D6E4",size=4):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcBorders')): tcPr.remove(old)
    b=OxmlElement('w:tcBorders')
    for side in('top','left','bottom','right'):
        el=OxmlElement(f'w:{side}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(size))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)
        b.append(el)
    tcPr.append(b)

def wcell(cell,text,bold=False,color=BLACK_RGB,align=WD_ALIGN_PARAGRAPH.LEFT,size=10.5):
    p=cell.paragraphs[0]; p.clear(); p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.name="Calibri"
    r.font.size=Pt(size); r.font.color.rgb=color

doc=Document()
for sec in doc.sections:
    sec.top_margin=Cm(2.54); sec.bottom_margin=Cm(2.54)
    sec.left_margin=Cm(2.54); sec.right_margin=Cm(2.54)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run("Bogota, 28 de mayo de 2026")
r.font.name="Calibri"; r.font.size=Pt(9.5); r.font.color.rgb=GRAY_RGB; r.italic=True

doc.add_paragraph("")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("CARGA TRANSPORTADA")
r.bold=True; r.font.name="Calibri"; r.font.size=Pt(16); r.font.color.rgb=NAVY_RGB

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("Corredor Ferreo Cienaga - La Dorada")
r.bold=True; r.font.name="Calibri"; r.font.size=Pt(13); r.font.color.rgb=NAVY_RGB

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("Periodo: Marzo - Mayo 2026  |  Unidad: Toneladas metricas (t)")
r.font.name="Calibri"; r.font.size=Pt(10); r.font.color.rgb=GRAY_RGB; r.italic=True

p=doc.add_paragraph()
p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(10)
pPr=p._p.get_or_add_pPr()
pBdr=OxmlElement('w:pBdr')
bot=OxmlElement('w:bottom')
bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'12')
bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'1A3A5C')
pBdr.append(bot); pPr.append(pBdr)

nrows=1+len(FILAS)+1
tabla=doc.add_table(rows=nrows,cols=6)
tabla.style='Table Grid'
for i,a in enumerate(ANCHOS_W):
    for row in tabla.rows: row.cells[i].width=Cm(a)

# Cabecera
hdr=tabla.rows[0]; hdr.height=Cm(0.9)
for j,txt in enumerate(CABECERAS):
    c=hdr.cells[j]; set_bg(c,NAVY_HEX); set_borders(c,color="1A3A5C",size=6)
    c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    aln=WD_ALIGN_PARAGRAPH.RIGHT if j>=2 else WD_ALIGN_PARAGRAPH.LEFT
    wcell(c,txt,bold=True,color=WHITE_RGB,align=aln,size=10)

# Datos
for idx,fila in enumerate(FILAS):
    row=tabla.rows[idx+1]
    bg=LGRAY_HEX if idx%2==1 else WHITE_HEX
    for j,txt in enumerate(fila):
        c=row.cells[j]; set_bg(c,bg); set_borders(c,color="C8D4E8",size=4)
        c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        aln=WD_ALIGN_PARAGRAPH.RIGHT if j>=2 else WD_ALIGN_PARAGRAPH.LEFT
        wcell(c,txt,align=aln,size=10.5)

# Total
row_tot=tabla.rows[nrows-1]; row_tot.height=Cm(0.85)
for j,txt in enumerate(TOTAL_ROW):
    c=row_tot.cells[j]; set_bg(c,LBLUE_HEX); set_borders(c,color="1A3A5C",size=6)
    c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    aln=WD_ALIGN_PARAGRAPH.RIGHT if j>=2 else WD_ALIGN_PARAGRAPH.LEFT
    wcell(c,txt,bold=True,color=NAVY_RGB,align=aln,size=10.5)

doc.add_paragraph("")
p=doc.add_paragraph()
p.paragraph_format.left_indent=Cm(0.3)
r=p.add_run(
    "Nota: Corresponde a 12 trenes intermodales operados en el periodo "
    "(6 en la direccion Cienaga -> La Dorada y 6 en La Dorada -> Cienaga). "
    "Las cifras incluyen toda la carga movilizada en ambos sentidos del corredor.")
r.font.name="Calibri"; r.font.size=Pt(9); r.font.color.rgb=GRAY_RGB; r.italic=True

OUT_DOCX = r"C:\Users\jony9\Downloads\Jony 11-25\Carga_Transportada_Mar-May2026_v2.docx"
doc.save(OUT_DOCX)
print(f"Word OK: {OUT_DOCX}")
